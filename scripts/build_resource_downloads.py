"""Build public downloads from shared page content. Development-only dependencies:
reportlab, python-docx and openpyxl. No runtime document generation or macros.
Run from any directory: python scripts/build_resource_downloads.py
"""
from pathlib import Path
from datetime import date
import json
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'static' / 'downloads'
OUT.mkdir(parents=True, exist_ok=True)
DATA = json.loads((ROOT / 'content' / 'provider-resources.json').read_text())
TEAL = '077A5C'
BASE = 'https://www.tuitionping.com'
FONT_DIR = Path('/usr/share/fonts/truetype/dejavu')
pdfmetrics.registerFont(TTFont('Resource', str(FONT_DIR / 'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('ResourceBold', str(FONT_DIR / 'DejaVuSans-Bold.ttf')))
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='ResourceTitle', fontName='ResourceBold', fontSize=23, leading=29, textColor=HexColor('#'+TEAL), spaceAfter=18))
styles.add(ParagraphStyle(name='ResourceHeading', fontName='ResourceBold', fontSize=12, leading=17, spaceBefore=15, spaceAfter=7, textColor=HexColor('#'+TEAL), keepWithNext=True))
styles.add(ParagraphStyle(name='ResourceBody', fontName='Resource', fontSize=10, leading=15, spaceAfter=9))


def documents(stem, title, sections, route):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.7)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Title'].font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_heading(title, 0)
    doc.add_paragraph('TuitionPing · Free provider resource · Updated October 6, 2026')
    flow = [Paragraph(escape(title), styles['ResourceTitle']), Paragraph('TuitionPing · Free provider resource · Updated October 6, 2026', styles['ResourceBody'])]
    for heading, paragraphs in sections:
        doc.add_heading(heading, 1)
        group = [Paragraph(escape(heading), styles['ResourceHeading'])]
        for text in paragraphs:
            paragraph = doc.add_paragraph(text)
            paragraph.paragraph_format.keep_with_next = text != paragraphs[-1]
            group.append(Paragraph(escape(text), styles['ResourceBody']))
        flow.append(KeepTogether(group))
    doc.add_paragraph('More free resources: '+BASE+route)
    doc.core_properties.title = title
    doc.core_properties.author = 'TuitionPing / Hirsch Commerce LLC'
    doc.save(OUT / (stem+'.docx'))
    flow += [Spacer(1, 10), Paragraph('More free resources: <link href="'+BASE+route+'">'+BASE+route+'</link>', styles['ResourceBody'])]
    def footer(canvas, document):
        canvas.setFont('Resource', 8)
        canvas.setFillColor(HexColor('#526376'))
        canvas.drawString(48, 28, 'TuitionPing · tuitionping.com/guides')
        canvas.drawRightString(564, 28, str(document.page))
    SimpleDocTemplate(str(OUT / (stem+'.pdf')), pagesize=(612,792), rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=48, title=title, author='TuitionPing / Hirsch Commerce LLC').build(flow, onFirstPage=footer, onLaterPages=footer)

policy_sections = [
    ('How to use this template', ['Replace all bracketed fields. This is a starting point, not a determination that a fee or care change is permitted. Check your enrollment agreement and applicable rules.']),
    ('Editable policy', DATA['policy']),
    *[(item['title'], [item['text']]) for item in DATA['examples']],
    ('Before you share it', ['State the deadline, calendar-day rule, one-time or recurring fee trigger, cap, payment methods, contact details, and notice process. Document arrangements and waivers. Do not treat a parent reporting PAID as verified receipt.'])]
documents('daycare-late-fee-policy', 'Daycare late-fee policy template', policy_sections, '/late-fee-policy')
reminder_sections = [('Before sending', ['Replace every placeholder. Check your balance and payment records first. Use the family’s preferred language, keep messages private, send with permission, and respect opt-outs. These are standalone templates; configure your account separately.'])]
for item in DATA['reminders']:
    reminder_sections.append((item['title'], [item['use'], 'English: '+item['en'], 'Español: '+item['es']]))
reminder_sections.append(('Using PAID with TuitionPing', ['PAID is the reply keyword in either language. A reply is parent-reported, not bank verification. The provider-verified receipt template is manual and should only be sent after checking the actual payment.']))
documents('daycare-tuition-reminders', 'Daycare tuition reminder texts', reminder_sections, '/guides/tuition-reminder-templates')

wb = Workbook()
info = wb.active
info.title = 'Start here'
info.append(['TuitionPing · Free daycare tuition payment tracker'])
info.append(['One row per bill. Weekly, every-two-week and monthly periods are supported in this standalone workbook.'])
info.append(['1. On Tracker, set B2 to your review date. It is fixed; update it each time you review balances.'])
info.append(['2. Enter an account label, billing period, tuition amount and due date. Avoid children’s names where possible.'])
info.append(['3. Enter assessed fees and approved credits. The sheet does not calculate or assess fees for you.'])
info.append(['4. Enter only verified payment amounts in column G. A parent report in H does not reduce the balance.'])
info.append(['5. Filter Status to Overdue; log your last contact and next action. Keep the completed workbook private.'])
info.append(['Amounts are USD. Balance = tuition + assessed fees − credits − verified payments, minimum zero.'])
info.append(['Formula columns I–K update in Excel or Google Sheets. No macros, messages, account sync or payment processing.'])
info.append(['Blank tuition is not classified. Blank due dates show Needs due date. A blank review date shows Set review date.'])
info.append(['Overpayments and multiple payment transactions need a separate ledger. This tracks the remaining balance per bill.'])
info.append(['Google Sheets: upload the .xlsx file to Drive, open with Google Sheets, save as a Google spreadsheet.'])
info.append(['After import, confirm the Worked example balance is $190 and overdue days are 7.'])
info.append(['This is not TuitionPing’s family-roster CSV import format.'])
info.append(['Get the calculator, editable policies and bilingual texts: '+BASE+'/guides'])
info.cell(15,1).hyperlink=BASE+'/guides'
info.column_dimensions['A'].width=110
for row in info:
    row[0].alignment=Alignment(wrap_text=True, vertical='top')
    info.row_dimensions[row[0].row].height=33
info.cell(1,1).font=Font(size=18,bold=True,color=TEAL)

headers = ['Family / account label','Billing period','Due date','Tuition ($)','Assessed fees ($)','Approved credits ($)','Verified payments ($)','Parent reported PAID / notes','Balance due ($)','Days overdue','Status','Last contact date','Next action']

def tracker_sheet(name, rows, review_date):
    ws=wb.create_sheet(name)
    ws.append([name+' · one row per tuition bill'])
    ws.append(['Review date',review_date]);ws['B2'].number_format='mmm d, yyyy'
    ws.append(['Total remaining balance', '=SUM(I6:I'+str(rows+5)+')']);ws['B3'].number_format='$#,##0.00'
    ws.append(['Edit blue input columns A–H and L–M. I–K contain formulas. See Start here before use.'])
    ws.append(headers)
    for n in range(6,rows+6):
        ws[f'I{n}']=f'=IF(OR(A{n}="",D{n}=""),"",MAX(0,SUM(D{n}:E{n})-SUM(F{n}:G{n})))'
        ws[f'J{n}']=f'=IF(OR(A{n}="",D{n}="",C{n}="",$B$2=""),"",IF(I{n}=0,0,MAX(0,$B$2-C{n})))'
        ws[f'K{n}']=f'=IF(OR(A{n}="",D{n}=""),"",IF($B$2="","Set review date",IF(I{n}=0,"Paid",IF(C{n}="","Needs due date",IF(C{n}<$B$2,"Overdue",IF(C{n}=$B$2,"Due today","Upcoming"))))))'
        for col in ['C','L']:ws[f'{col}{n}'].number_format='mmm d, yyyy'
        for col in ['D','E','F','G','I']:ws[f'{col}{n}'].number_format='$#,##0.00'
        for c in ws[n]:
            c.alignment=Alignment(vertical='top',wrap_text=True)
            c.fill=PatternFill('solid',fgColor='EAF4FC' if c.column in [1,2,3,4,5,6,7,8,12,13] else 'F1F5F9')
    table=Table(displayName='Bills'+name.replace(' ','').replace('-',''),ref=f'A5:M{rows+5}')
    table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
    ws.add_table(table)
    for c in ws[5]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=TEAL);c.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[5].height=40
    ws.freeze_panes='D6'
    widths=[25,22,17,16,19,21,23,32,21,17,20,20,38]
    for i,width in enumerate(widths,1):ws.column_dimensions[ws.cell(5,i).column_letter].width=width
    ws.sheet_view.zoomScale=80
    ws['A1'].font=Font(bold=True,size=16,color=TEAL)
    dv=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=0,allow_blank=True)
    dv.errorTitle='Use a nonnegative amount';dv.error='Enter a number of zero or more.';dv.showErrorMessage=True;dv.errorStyle='stop';ws.add_data_validation(dv);dv.add(f'D6:G{rows+5}')
    dates=DataValidation(type='date',operator='between',formula1='DATE(1900,1,1)',formula2='DATE(9999,12,31)',allow_blank=True)
    dates.showErrorMessage=True;dates.errorStyle='stop';dates.error='Enter a valid date.';ws.add_data_validation(dates);dates.add('B2');dates.add(f'C6:C{rows+5}');dates.add(f'L6:L{rows+5}')
    ws.conditional_formatting.add(f'K6:K{rows+5}',FormulaRule(formula=['K6="Overdue"'],fill=PatternFill('solid',fgColor='FEE2E2')))
    ws.print_title_rows='1:5';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    return ws

tracker_sheet('Tracker',200,date(2026,10,6))
example=tracker_sheet('Worked example',4,date(2026,10,12))
for col,value in enumerate(['Sample family A','Week of Oct 5',date(2026,10,5),300,10,20,100,'Parent reported PAID; not yet verified'],1):example.cell(6,col,value)
for col,value in enumerate(['Sample family B','October',date(2026,10,1),1000,0,0,1000,'Payment checked by provider'],1):example.cell(7,col,value)
example['M6']='Verify report; remaining balance $190, 7 days overdue.'
example['M7']='Paid: balance $0; overdue days 0.'
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
wb.save(OUT / 'daycare-tuition-payment-tracker.xlsx')
readme='''TUITIONPING — FREE TUITION COLLECTION KIT
Updated October 6, 2026

Start with the editable policy, customize the bracketed fields, and check your agreement and applicable rules.
Use the English/Spanish reminder templates for private, permission-based communication.
Open the payment tracker and read Start here. Set the review date before entering bills.
Only verified payments reduce the spreadsheet balance; a parent's PAID report alone does not.
Word documents are editable; PDFs are printable references. The tracker has no macros.
The standalone workbook is not a TuitionPing CSV import or a payment processor.

These free materials are yours to adapt for your program. For current versions and to recommend the resource to others:
https://www.tuitionping.com/guides
Calculator: https://www.tuitionping.com/tools/late-fee-calculator
Support: https://www.tuitionping.com/support
'''
(OUT/'START-HERE.txt').write_text(readme)
with zipfile.ZipFile(OUT/'daycare-tuition-collection-kit.zip','w',zipfile.ZIP_DEFLATED) as z:
    for f in sorted(OUT.iterdir()):
        if f.suffix in ['.pdf','.docx','.xlsx','.txt']:z.write(f,f.name)
print('Built', [(f.name,f.stat().st_size) for f in sorted(OUT.iterdir())])
